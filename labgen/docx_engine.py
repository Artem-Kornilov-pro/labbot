"""Движок отчёта .docx по методичке: только родные объекты Word (OMML-формулы, фигуры, поле TOC).

Общий для всех лабораторных. Конкретная лаба собирает отчёт вызовами методов Report.
"""
import io
import re
from dataclasses import dataclass, field

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, Mm, Cm, RGBColor

from .textmetrics import text_width_pt

M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
WPG = "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"
NS = f'xmlns:w="{W}" xmlns:m="{M}" xmlns:a="{A}" xmlns:wp="{WP}" xmlns:wps="{WPS}" xmlns:wpg="{WPG}"'
FONT = "Times New Roman"
MONO = "Courier New"
C = WD_ALIGN_PARAGRAPH.CENTER
J = WD_ALIGN_PARAGRAPH.JUSTIFY
CONTENT_W_MM = 165


def visible(xml):
    """Видимый текст формулы (без разметки) – для контроля ширины строки."""
    import html
    return html.unescape(re.sub(r"<[^>]+>", "", xml))


def width(text):
    """Условная ширина строки формулы в символах: пробел – половина символа."""
    return sum(0.5 if ch == " " else 1 for ch in text)


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def X(xml):
    return parse_xml(xml.replace("<NS>", NS))


# ================================================================ формулы Word (OMML)
MATH_SZ = 24   # 12 pt
MATH_MAX = 66  # максимальная видимая длина строки формулы (символов), см. tests


def _mrun(text, plain=False):
    rpr = '<m:rPr><m:sty m:val="p"/></m:rPr>' if plain else ""
    return (f'<m:r>{rpr}<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/>'
            f'<w:sz w:val="{MATH_SZ}"/><w:szCs w:val="{MATH_SZ}"/></w:rPr>'
            f'<m:t xml:space="preserve">{esc(text)}</m:t></m:r>')


_FUNCS = r"sin|cos|max|min"


def m(text):
    """Строка формулы: кириллица и имена функций — прямым шрифтом, остальное как в редакторе формул Word."""
    text = text.replace("\\", "∖")       # целочисленное деление (в OMML «\\» ломается в LibreOffice)
    out = []
    pat = rf"[А-Яа-яЁё]+(?:\s+[А-Яа-яЁё]+)*|{_FUNCS}|[^А-Яа-яЁё]+?(?=[А-Яа-яЁё]|{_FUNCS}|$)"
    for tok in re.findall(pat, text):
        out.append(_mrun(tok, plain=bool(re.match(rf"[А-Яа-яЁё]|{_FUNCS}", tok))))
    return "".join(out)


def sup(base, exp):
    return f'<m:sSup><m:e>{m(base)}</m:e><m:sup>{m(exp)}</m:sup></m:sSup>'


def bar(inner):
    return f'<m:bar><m:barPr><m:pos m:val="top"/></m:barPr><m:e>{inner}</m:e></m:bar>'


def frac(num, den):
    return f'<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>'


def nary(ch, sub, sup_, body):
    return (f'<m:nary><m:naryPr><m:chr m:val="{ch}"/><m:limLoc m:val="undOvr"/></m:naryPr>'
            f'<m:sub>{sub}</m:sub><m:sup>{sup_}</m:sup><m:e>{body}</m:e></m:nary>')


def nsum(sub, sup_, body):
    return nary("∑", sub, sup_, body)


def nprod(sub, sup_, body):
    return nary("∏", sub, sup_, body)


def brace(rows):
    """Левая фигурная скобка на все строки (как в методичке); строки выровнены по левому краю."""
    mrs = "".join(f"<m:mr><m:e>{r}</m:e></m:mr>" for r in rows)
    return ('<m:d><m:dPr><m:begChr m:val="{"/><m:endChr m:val=""/></m:dPr><m:e>'
            '<m:m><m:mPr><m:baseJc m:val="center"/><m:rSpRule m:val="1"/><m:mcs><m:mc><m:mcPr>'
            '<m:count m:val="1"/><m:mcJc m:val="left"/></m:mcPr></m:mc></m:mcs></m:mPr>'
            f'{mrs}</m:m></m:e></m:d>')


def rng(a, b):
    """Диапазон с чертой сверху: a, b̄."""
    return bar(m(f"{a}, {b}"))


def loop(var, a, b):
    """«для i = 1,n» с чертой над диапазоном."""
    return m(f"для {var} = ") + rng(a, b)


def F(s):
    """Короткая запись строки формулы: {a..b} -> диапазон с чертой. Пример: F('∀ i = {1..n}: R[i] > 0')."""
    out, pos = [], 0
    for mt in re.finditer(r"\{([^{}]+?)\.\.([^{}]+?)\}", s):
        if mt.start() > pos:
            out.append(m(s[pos:mt.start()]))
        out.append(rng(mt.group(1).strip(), mt.group(2).strip()))
        pos = mt.end()
    if pos < len(s):
        out.append(m(s[pos:]))
    return "".join(out)


# ================================================================ спецификация (фигуры Word)
EMU = 12700
SPEC_SZ = 12
SPEC_LH = 14


def _wrap(line, width_pt):
    out, cur = [], ""
    for w in line.split(" "):
        t = (cur + " " + w) if cur else w
        if text_width_pt(t, SPEC_SZ) <= width_pt or not cur:
            cur = t
        else:
            out.append(cur)
            cur = w
    out.append(cur)
    return out


class _Shapes:
    def __init__(self, start_id):
        self.sid = start_id

    def next(self):
        self.sid += 1
        return self.sid

    @staticmethod
    def _xfrm(x, y, w, h):
        return (f'<a:xfrm><a:off x="{int(x*EMU)}" y="{int(y*EMU)}"/>'
                f'<a:ext cx="{int(w*EMU)}" cy="{int(h*EMU)}"/></a:xfrm>')

    def text(self, x, y, w, h, lines, border=True, mono=False):
        sid = self.next()
        ln = ('<a:ln w="9525"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln>' if border
              else '<a:ln><a:noFill/></a:ln>')
        font = MONO if mono else FONT
        size = SPEC_SZ - 1 if mono else SPEC_SZ
        paras = "".join(
            f'<w:p><w:pPr><w:spacing w:before="0" w:after="0" w:line="{SPEC_LH*20}" w:lineRule="exact"/></w:pPr>'
            f'<w:r><w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:cs="{font}"/>'
            f'<w:sz w:val="{size*2}"/><w:szCs w:val="{size*2}"/></w:rPr>'
            f'<w:t xml:space="preserve">{esc(s)}</w:t></w:r></w:p>' for s in lines)
        return (f'<wps:wsp><wps:cNvPr id="{sid}" name="Надпись {sid}"/><wps:cNvSpPr txBox="1"/>'
                f'<wps:spPr>{self._xfrm(x, y, w, h)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
                f'<a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>{ln}</wps:spPr>'
                f'<wps:txbx><w:txbxContent>{paras}</w:txbxContent></wps:txbx>'
                f'<wps:bodyPr rot="0" vert="horz" wrap="square" lIns="63500" tIns="38100" rIns="63500" '
                f'bIns="25400" anchor="t" anchorCtr="0"><a:noAutofit/></wps:bodyPr></wps:wsp>')

    def geom(self, x, y, w, h, prst, fill=False):
        sid = self.next()
        fill_xml = '<a:solidFill><a:srgbClr val="000000"/></a:solidFill>' if fill else '<a:noFill/>'
        return (f'<wps:wsp><wps:cNvPr id="{sid}" name="Фигура {sid}"/><wps:cNvSpPr/>'
                f'<wps:spPr>{self._xfrm(x, y, w, h)}<a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom>{fill_xml}'
                f'<a:ln w="12700"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln></wps:spPr>'
                f'<wps:bodyPr/></wps:wsp>')


# ================================================================ отчёт
@dataclass
class TitleInfo:
    author: str
    group: str = ""
    teacher: str = ""
    variant: str = ""            # «№15 (16, 6)»
    lab_no: int = 1
    year: str = "2026"
    dept: str = "Департамент компьютерной инженерии"
    grade_cols: list = field(default_factory=lambda: ["Итог.\nоценка"])


SECTIONS_DEFAULT = ["Задание", "Постановка задачи", "Метод решения задачи", "Внешняя спецификация",
                    "Описание алгоритма на псевдокоде", "Листинг программы", "Тесты"]


class Report:
    def __init__(self, title: TitleInfo, sections=None):
        self.t = title
        self.sections = sections or SECTIONS_DEFAULT
        self.bm = {name: f"_Toc{1000000 + i:09d}" for i, name in enumerate(self.sections)}
        self._bm_id = 0
        self._shape_id = 1000
        self.doc = Document()
        self.body = self.doc.element.body
        self._styles()
        self._page()
        self.lines_est = {}           # оценка высоты разделов для кэша оглавления
        self.math_widths = []         # (видимая длина, текст) каждой строки формулы
        self._cur = None

    # ------------------------------------------------------------ стили/страница
    def _get_style(self, name):
        try:
            return self.doc.styles[name]
        except KeyError:
            return self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

    def _styles(self):
        doc = self.doc
        st = doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(14)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.line_spacing = 1.15

        def heading(style, size, center, before, after, pbb):
            s = doc.styles[style]
            s.font.name = FONT
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.italic = False
            s.font.color.rgb = RGBColor(0, 0, 0)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                s.element.rPr.rFonts.attrib.pop(qn(a), None)
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                s.element.rPr.rFonts.set(qn(a), FONT)
            pf = s.paragraph_format
            if center:
                pf.alignment = C
            pf.space_before, pf.space_after = Pt(before), Pt(after)
            pf.page_break_before = pbb
            pf.keep_with_next = True

        heading("Heading 1", 16, True, 0, 12, True)
        heading("Heading 2", 14, False, 12, 6, False)

        code = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = st
        code.font.name = MONO
        code.font.size = Pt(8.5)
        code.element.rPr.rFonts.set(qn("w:eastAsia"), MONO)
        code.element.rPr.rFonts.set(qn("w:cs"), MONO)
        code.paragraph_format.line_spacing = 1.0

        toc1 = self._get_style("toc 1")
        toc1.base_style = st
        toc1.paragraph_format.space_after = Pt(6)
        toc1.paragraph_format.tab_stops.add_tab_stop(Mm(CONTENT_W_MM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        toch = self._get_style("TOC Heading")
        if toch.element.rPr is not None and toch.element.rPr.rFonts is not None:
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                toch.element.rPr.rFonts.attrib.pop(qn(a), None)
        toch.base_style = st
        toch.font.name = FONT
        toch.font.size = Pt(16)
        toch.font.bold = True
        toch.font.color.rgb = RGBColor(0, 0, 0)
        toch.paragraph_format.alignment = C
        toch.paragraph_format.space_after = Pt(12)
        toch.paragraph_format.page_break_before = True
        toch.element.get_or_add_pPr().append(X('<w:outlineLvl <NS> w:val="9"/>'))

    def _page(self):
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin, sec.right_margin = Mm(30), Mm(15)
        sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)
        sec.different_first_page_header_footer = True
        sec._sectPr.append(X('<w:pgNumType <NS> w:start="0"/>'))
        fp = sec.footer.paragraphs[0]
        fp.alignment = C
        for xml in ('<w:fldChar w:fldCharType="begin"/>',
                    '<w:instrText xml:space="preserve"> PAGE </w:instrText>',
                    '<w:fldChar w:fldCharType="separate"/>', '<w:t>1</w:t>',
                    '<w:fldChar w:fldCharType="end"/>'):
            fp._p.append(X(f'<w:r <NS>><w:rPr><w:sz w:val="24"/></w:rPr>{xml}</w:r>'))
        ffp = sec.first_page_footer.paragraphs[0]
        ffp.alignment = C
        r = ffp.add_run(f"МОСКВА {self.t.year}")
        r.bold = True
        r.font.size = Pt(14)

    # ------------------------------------------------------------ базовые элементы
    def _est(self, lines):
        if self._cur:
            self.lines_est[self._cur] = self.lines_est.get(self._cur, 0) + lines

    @staticmethod
    def add_runs(p, text, bold=False, size=None, italic=False, underline=False, font=None):
        """Текст с разметкой: _слово_ — подчёркнуто, *слово* — жирно."""
        for part in re.split(r"(_[^_\s][^_]*_|\*[^*]+\*)", text):
            if not part:
                continue
            u, b = underline, bold
            if part.startswith("_") and part.endswith("_") and len(part) > 2:
                part, u = part[1:-1], True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                part, b = part[1:-1], True
            r = p.add_run(part)
            r.bold, r.italic, r.underline = b, italic, u
            if size:
                r.font.size = Pt(size)
            if font:
                r.font.name = font
                r._r.rPr.rFonts.set(qn("w:eastAsia"), font)
        return p

    def para(self, text="", bold=False, align=None, size=None, before=0, after=0, indent=None,
             keep=False, raw=False):
        p = self.doc.add_paragraph()
        if text:
            if raw:
                p.add_run(text)
            else:
                self.add_runs(p, text, bold=bold, size=size)
        pf = p.paragraph_format
        if align is not None:
            pf.alignment = align
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        if indent is not None:
            pf.left_indent = Cm(indent)
        if keep:
            pf.keep_with_next = True
        chars = int((CONTENT_W_MM - (indent or 0) * 10) / 165 * 78 * 14 / (size or 14))
        self._est(max(1, -(-len(text) // max(chars, 20))) + (before + after) / 16)
        return p

    def heading(self, text):
        """Заголовок 1-го уровня (с новой страницы) с закладкой для оглавления."""
        h = self.doc.add_heading(text, level=1)
        self._bm_id += 1
        h._p.insert(1, X(f'<w:bookmarkStart <NS> w:id="{self._bm_id}" w:name="{self.bm[text]}"/>'))
        h._p.append(X(f'<w:bookmarkEnd <NS> w:id="{self._bm_id}"/>'))
        self._cur = text
        self._est(2)
        return h

    def heading2(self, text):
        self._est(2)
        return self.doc.add_heading(text, level=2)

    def math_inline(self, p, inner):
        p._p.append(X(f'<m:oMath <NS>>{inner}</m:oMath>'))

    def fline(self, prefix, inner, indent=None, after=2):
        """Абзац: обычный текст + формула в строке."""
        txt = prefix + visible(inner)
        self.math_widths.append((width(txt) + (indent or 0) * 3, txt))
        p = self.para(prefix, indent=indent, after=after)
        self.math_inline(p, inner)
        return p

    def math_block(self, inner, after=6, indent=0, rows=1):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.left_indent = Cm(indent)
        p._p.append(X(f'<m:oMathPara <NS>><m:oMathParaPr><m:jc m:val="left"/></m:oMathParaPr>'
                      f'<m:oMath>{inner}</m:oMath></m:oMathPara>'))
        self._est(rows * 1.3 + after / 16)
        return p

    def method(self, items, after=10):
        """items: str — строка формулы; list — блок под фигурной скобкой (строки или вложенные list)."""
        def conv(it, depth=0):
            if isinstance(it, list):
                return brace([conv(x, depth + 1) for x in it])
            self.math_widths.append((width(visible(F(it))) + 2 * depth, it))
            return F(it)

        def height(it):
            return sum(height(x) for x in it) if isinstance(it, list) else 1

        for i, it in enumerate(items):
            last = i == len(items) - 1
            self.math_block(conv(it), after=(after if last else 1), rows=height(it))

    # ------------------------------------------------------------ таблицы
    @staticmethod
    def _cell_borders(cell, val="single"):
        tcPr = cell._tc.get_or_add_tcPr()
        xml = "".join(f'<w:{s} w:val="{val}" w:sz="4" w:space="0" w:color="000000"/>'
                      for s in ("top", "left", "bottom", "right"))
        tcPr.append(X(f'<w:tcBorders <NS>>{xml}</w:tcBorders>'))

    @staticmethod
    def _table_borders(tbl, val="single"):
        xml = "".join(f'<w:{s} w:val="{val}" w:sz="4" w:space="0" w:color="000000"/>'
                      for s in ("top", "left", "bottom", "right", "insideH", "insideV"))
        tbl._tbl.tblPr.append(X(f'<w:tblBorders <NS>>{xml}</w:tblBorders>'))

    @staticmethod
    def _set_widths(tbl, widths):
        tbl.autofit = False
        tbl._tbl.tblPr.append(X('<w:tblLayout <NS> w:type="fixed"/>'))
        for row in tbl.rows:
            for c, w in zip(row.cells, widths):
                c.width = w
        for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w.twips)))

    def _cell_text(self, cell, lines, size=12, bold=False, align=None, font=None, spacing=1.0, style=None):
        cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
        for ln in lines:
            p = cell.add_paragraph(style=style)
            p.paragraph_format.line_spacing = spacing
            if align is not None:
                p.alignment = align
            if style:
                p.add_run(ln)
            else:
                self.add_runs(p, ln, bold=bold, size=size, font=font)

    # ------------------------------------------------------------ титульный лист и содержание
    def title_page(self):  # noqa: C901
        t = self.t
        self.para("Федеральное государственное автономное образовательное учреждение высшего образования "
                  "«Национальный исследовательский университет «Высшая школа экономики»", bold=True, align=C)
        self.para("Московский институт электроники и математики им. А.Н. Тихонова НИУ ВШЭ",
                  bold=True, align=C, before=18)
        self.para(t.dept, bold=True, align=C, before=18)
        self.para("Курс: Алгоритмизация и программирование", bold=True, align=C, before=36)
        self.para("ОТЧЕТ", bold=True, align=C, before=48)
        self.para(f"по лабораторной работе №{t.lab_no}", bold=True, align=C, before=6, after=36)

        gcols = t.grade_cols
        rows = [("Раздел", "Мах\nоценка", *gcols), ("Постановка", "0,5"), ("Метод", "1"),
                ("Спецификация", "0,5"), ("Алгоритм", "1,5"), ("Работа программы", "1"),
                ("Листинг", "0,5"), ("Тесты", "1"), ("Вопросы", "2"), ("Доп. задание", "2")]
        ng = 2 + len(gcols)
        gw = [Mm(31), Mm(17)] + [Mm(17)] * len(gcols)
        used = sum(w.mm for w in gw)
        widths = gw + [Mm(5), Mm(CONTENT_W_MM - used - 5)]
        tbl = self.doc.add_table(rows=len(rows), cols=ng + 2)
        self._table_borders(tbl, "nil")
        self._set_widths(tbl, widths)
        for r, vals in zip(tbl.rows, rows):
            r.height = Mm(8)
            vals = list(vals) + [""] * (ng - len(vals))
            for c, v in zip(r.cells[:ng], vals):
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                self._cell_borders(c)
                self._cell_text(c, v.split("\n"), size=11 if r is not tbl.rows[0] else 10,
                                bold=(r is tbl.rows[0]))
        info = tbl.cell(0, ng + 1).merge(tbl.cell(len(rows) - 1, ng + 1))
        for rr in range(len(rows)):
            tbl.cell(rr, ng).text = ""
        info.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        blank = "________________"
        self._cell_text(info, [f"Студент: {t.author}", f"Группа: {t.group or blank}",
                               f"Вариант: {t.variant}", f"Руководитель: {t.teacher or blank}",
                               "Оценка: ________________", "Дата сдачи: _____________"],
                        size=14, bold=True, spacing=1.15)
        for p in info.paragraphs:
            p.paragraph_format.space_after = Pt(10)

    def toc_placeholder(self):
        """Вставить поле оглавления; номера страниц заполняются в finalize()."""
        self._toc_anchor = self.doc.add_paragraph()

    def _toc_xml(self, pages):
        def r_(inner):
            return f"<w:r>{inner}</w:r>"
        toc_style = self.doc.styles["toc 1"].style_id
        toch_style = self.doc.styles["TOC Heading"].style_id
        entries = []
        for i, name in enumerate(self.sections):
            start = ""
            if i == 0:
                start = (r_('<w:fldChar w:fldCharType="begin"/>')
                         + r_('<w:instrText xml:space="preserve"> TOC \\o "1-1" \\h \\z \\u </w:instrText>')
                         + r_('<w:fldChar w:fldCharType="separate"/>'))
            entries.append(
                f'<w:p><w:pPr><w:pStyle w:val="{toc_style}"/></w:pPr>{start}'
                f'<w:hyperlink w:anchor="{self.bm[name]}" w:history="1">'
                f'{r_(f"<w:t>{esc(name)}</w:t>")}{r_("<w:tab/>")}'
                + r_('<w:fldChar w:fldCharType="begin"/>')
                + r_(f'<w:instrText xml:space="preserve"> PAGEREF {self.bm[name]} \\h </w:instrText>')
                + r_('<w:fldChar w:fldCharType="separate"/>')
                + r_(f"<w:t>{pages.get(name, '')}</w:t>")
                + r_('<w:fldChar w:fldCharType="end"/>')
                + '</w:hyperlink></w:p>')
        return X('<w:sdt <NS>><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/>'
                 '<w:docPartUnique/></w:docPartObj></w:sdtPr><w:sdtContent>'
                 f'<w:p><w:pPr><w:pStyle w:val="{toch_style}"/></w:pPr><w:r><w:t>Содержание</w:t></w:r></w:p>'
                 + "".join(entries) +
                 '<w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p></w:sdtContent></w:sdt>')

    # ------------------------------------------------------------ спецификация
    SPEC_MAX_MM = 225

    def spec(self, items):
        """items: ("box", [строки], скобка_слева) | ("label", текст) | ("rep_start",) | ("rep_end", условие)
        | ("mono", [строки]) — окно с моноширинным текстом (матрица).
        Длинная спецификация делится на несколько групп (группа не переносится между страницами)."""
        chunk, chunks, depth = [], [], 0
        for i, it in enumerate(items):
            boundary = depth == 0 and it[0] in ("box", "mono", "label", "rep_start") \
                and not (i > 0 and items[i - 1][0] == "label")
            if boundary and chunk and self._spec_render(chunk + [it], dry=True) > self.SPEC_MAX_MM:
                chunks.append(chunk)
                chunk = []
            chunk.append(it)
            depth += it[0] == "rep_start"
            depth -= it[0] == "rep_end"
        chunks.append(chunk)
        for ch in chunks:
            self._spec_render(ch)

    def _spec_render(self, items, dry=False):
        sh = _Shapes(self._shape_id)
        total_w = CONTENT_W_MM / 25.4 * 72
        X_OUT, BR = 0, 7
        X_IN = X_OUT + BR + 4
        X_BOX = X_IN + BR + 4
        X_R = total_w - BR - 18
        X_RBR = X_R + 4
        GAP, LABEL_H = 6, 20
        inner_w = X_R - X_BOX - 16
        shapes, stack, y = [], [], 2
        for it in items:
            kind = it[0]
            if kind in ("box", "mono"):
                mono = kind == "mono"
                lines = list(it[1]) if mono else [s for ln in it[1] for s in _wrap(ln, inner_w)]
                h = len(lines) * SPEC_LH + 9
                shapes.append(sh.text(X_BOX, y, X_R - X_BOX, h, lines, mono=mono))
                if len(it) > 2 and it[2]:
                    shapes.append(sh.geom(X_IN, y, BR, h, "leftBrace"))
                y += h + GAP
            elif kind == "label":
                shapes.append(sh.text(X_BOX - 5, y, X_R - X_BOX, LABEL_H, [it[1]], border=False))
                y += LABEL_H
            elif kind == "rep_start":
                stack.append(y)
            elif kind == "rep_end":
                y0 = stack.pop()
                y1 = y - GAP
                shapes.append(sh.geom(X_OUT, y0, BR, y1 - y0, "leftBrace"))
                shapes.append(sh.geom(X_RBR, y0, BR, y1 - y0, "rightBrace"))
                shapes.append(sh.geom(X_RBR + BR + 3, y0 - 2, 9, 9, "star5", fill=True))
                shapes.append(sh.text(X_BOX - 5, y, X_R - X_BOX, LABEL_H, [it[1]], border=False))
                y += LABEL_H + GAP
        H = y + 2
        if dry:
            return H / 72 * 25.4
        cx, cy = int(total_w * EMU), int(H * EMU)
        xml = (f'<w:p <NS>><w:pPr><w:spacing w:before="0" w:after="120" w:line="240" w:lineRule="auto"/></w:pPr>'
               f'<w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
               f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
               f'<wp:docPr id="{sh.next()}" name="Внешняя спецификация"/><wp:cNvGraphicFramePr/>'
               f'<a:graphic><a:graphicData uri="{WPG}"><wpg:wgp><wpg:cNvGrpSpPr/>'
               f'<wpg:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/>'
               f'<a:chOff x="0" y="0"/><a:chExt cx="{cx}" cy="{cy}"/></a:xfrm></wpg:grpSpPr>'
               + "".join(shapes) +
               '</wpg:wgp></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>')
        self._shape_id = sh.sid
        self.body.insert(len(self.body) - 1, X(xml))
        self._est(H / 72 * 25.4 / 5.7 + 1)
        return H / 72 * 25.4

    # ------------------------------------------------------------ псевдокод, листинг, тесты
    def pseudocode(self, text):
        for line in text.strip("\n").split("\n"):
            lvl = len(line) - len(line.lstrip("\t"))
            p = self.para(line.strip("\t"), indent=1.25 * lvl)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(1)

    def listing(self, src):
        for line in src.rstrip("\n").split("\n"):
            self.doc.add_paragraph(line.replace("\t", "    "), style="Code")
        self._est(len(src.split("\n")) * 0.56)

    def tests_table(self, rows):
        """rows: [(исходные данные: list[str], результат: list[str], случай: str)]."""
        tbl = self.doc.add_table(rows=1, cols=4)
        self._table_borders(tbl)
        widths = [Mm(9), Mm(46), Mm(76), Mm(34)]
        for c, t in zip(tbl.rows[0].cells, ["№", "Исходные данные", "Результат", "Проверяемый случай"]):
            self._cell_text(c, [t], size=12, bold=True, align=C)
        for i, (inp, res, case) in enumerate(rows, 1):
            cells = tbl.add_row().cells
            self._cell_text(cells[0], [str(i)], size=12, align=C)
            self._cell_text(cells[1], inp, style="Code")
            self._cell_text(cells[2], res, style="Code")
            self._cell_text(cells[3], [case], size=11)
            self._est(max(len(inp), len(res)) * 0.6 + 0.5)
        self._set_widths(tbl, widths)
        tbl.rows[0]._tr.get_or_add_trPr().append(X('<w:tblHeader <NS>/>'))
        self.para("", after=6)

    def transcript(self, caption, text):
        self.para(caption, bold=True, before=8, after=2, keep=True)
        t = self.doc.add_table(rows=1, cols=1)
        self._table_borders(t)
        self._set_widths(t, [Mm(CONTENT_W_MM)])
        cell = t.rows[0].cells[0]
        cell._tc.get_or_add_tcPr().append(X('<w:shd <NS> w:val="clear" w:color="auto" w:fill="F2F2F2"/>'))
        lines = text.rstrip("\n").split("\n")
        self._cell_text(cell, lines, style="Code")
        t.rows[0]._tr.get_or_add_trPr().append(X('<w:cantSplit <NS>/>'))
        self._est(len(lines) * 0.56 + 1.5)

    # ------------------------------------------------------------ финал
    def _estimate_pages(self):
        per_page = 42.0            # строк 14 pt на странице
        pages, cur = {}, 2         # титульный = 0, содержание = 1
        for name in self.sections:
            pages[name] = cur
            n = self.lines_est.get(name, 1)
            cur += max(1, -(-int(n) // int(per_page)))
        return pages

    def _schema_sort(self, root):
        order = {
            "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText "
                    "vAlign hideMark",
            "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc "
                     "tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
            "sectPr": "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders "
                      "lnNumType pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter "
                      "docGrid printerSettings",
            "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr "
                   "shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE "
                   "autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents "
                   "suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr "
                   "sectPr pPrChange",
        }
        for tag, ordr in order.items():
            rank = {qn("w:" + t): i for i, t in enumerate(ordr.split())}
            for el in root.iter(qn("w:" + tag)):
                seen = {}
                for k in el:
                    if k.tag in rank:
                        seen[k.tag] = k
                kids = [k for k in el if k.tag not in rank or seen.get(k.tag) is k]
                for k in list(el):
                    el.remove(k)
                kids.sort(key=lambda k: rank.get(k.tag, len(rank)))
                for k in kids:
                    el.append(k)

    def finalize(self) -> bytes:
        pages = self._estimate_pages()
        anchor = self._toc_anchor._p
        anchor.addprevious(self._toc_xml(pages))
        anchor.getparent().remove(anchor)
        self._schema_sort(self.doc.element)
        self._schema_sort(self.doc.styles.element)
        settings = self.doc.settings.element
        z = settings.find(qn("w:zoom"))
        if z is not None:
            z.set(qn("w:percent"), "100")
        # Word обновит оглавление при открытии (номера страниц станут точными)
        settings.append(X('<w:updateFields <NS> w:val="true"/>'))
        self._schema_sort_settings(settings)
        self.doc.core_properties.author = self.t.author
        self.doc.core_properties.title = f"Отчёт по лабораторной работе №{self.t.lab_no}"
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()

    @staticmethod
    def _schema_sort_settings(settings):
        """updateFields должен стоять до элементов, идущих после него по схеме CT_Settings."""
        uf = settings.find(qn("w:updateFields"))
        after = ["hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
                 "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                 "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType",
                 "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"]
        settings.remove(uf)
        for k in settings:
            name = k.tag.split("}")[1]
            if name in after:
                k.addprevious(uf)
                return
        settings.append(uf)
